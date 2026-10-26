"""Generate the editable beginner guide with the bundled python-docx runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.8)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 25), ('Heading 1', 17), ('Heading 2', 13)]:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name != 'Title'
    style.paragraph_format.space_before = Pt(13 if name != 'Title' else 0)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.keep_with_next = True
code_style = doc.styles.add_style('Code block', 1)
code_style.font.name = 'Consolas'
code_style.font.size = Pt(9)
code_style.paragraph_format.space_after = Pt(3)
code_style.paragraph_format.line_spacing = 1.0

in_code = False
buffer = []
def flush():
    if buffer:
        doc.add_paragraph(' '.join(buffer))
        buffer.clear()

for line in (root / 'project_guide.md').read_text(encoding='utf-8').splitlines():
    if line.startswith('```'):
        flush()
        in_code = not in_code
    elif in_code:
        doc.add_paragraph(line, 'Code block')
    elif line == '<!-- page -->':
        flush()
        doc.add_page_break()
    elif line.startswith('# '):
        flush()
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        flush()
        doc.add_paragraph(line[3:], 'Heading 1')
    elif line.startswith('### '):
        flush()
        doc.add_paragraph(line[4:], 'Heading 2')
    elif line.startswith('- '):
        flush()
        doc.add_paragraph(line[2:], 'List Bullet')
    elif not line.strip():
        flush()
    else:
        buffer.append(line)
flush()

# Page numbering is useful for a printed lab guide.
footer = section.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run('RISC V project guide   |   ')
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Understanding the RISC V project files'
doc.core_properties.subject = 'Beginner guide to project files and block testing'
doc.core_properties.author = 'RISC-V SoC project'
# The bundled Word template may carry a decorative border on Title.
for element in (doc.styles.element, doc.element):
    for border in list(element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
out = root / 'project_guide.docx'
doc.save(out)
print(out)
